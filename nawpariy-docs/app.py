"""
Servicio de extracción de texto de documentos.

Existe porque el hosting compartido donde vive la API no puede hacer esto: tiene
shell_exec, exec y system deshabilitados, y ni siquiera ZipArchive compilado. Sin
una máquina con herramientas de verdad no hay forma de leer un Excel, un Word ni,
sobre todo, la foto de un cuaderno de campo.

No es un agente: no decide nada ni toca la base de datos. Recibe un archivo,
devuelve su texto y se olvida. Toda la inteligencia y todos los permisos siguen
del lado de la aplicación.
"""
import os
import glob
import shutil
import subprocess
import tempfile

from flask import Flask, request, jsonify

TOKEN = os.environ.get("DOC_TOKEN", "")
MAX_PAGINAS_OCR = int(os.environ.get("MAX_PAGINAS_OCR", "25"))
IDIOMA_OCR = os.environ.get("IDIOMA_OCR", "spa")

app = Flask(__name__)


def ejecutar(cmd, timeout=240, cwd=None):
    return subprocess.run(cmd, capture_output=True, timeout=timeout, cwd=cwd)


def texto_de_pdf(ruta, carpeta):
    """Primero la capa de texto; si el PDF es un escaneo, se pasa a OCR."""
    salida = os.path.join(carpeta, "pdf.txt")
    ejecutar(["pdftotext", "-layout", ruta, salida])

    texto = ""
    if os.path.exists(salida):
        with open(salida, "r", encoding="utf-8", errors="replace") as fh:
            texto = fh.read()

    if len(texto.strip()) >= 60:
        return texto, "capa de texto del PDF"

    # Sin texto aprovechable: es un escaneo. Se rasteriza y se pasa por OCR.
    ejecutar([
        "pdftoppm", "-r", "200", "-png",
        "-l", str(MAX_PAGINAS_OCR),
        ruta, os.path.join(carpeta, "pag"),
    ], timeout=600)

    paginas = sorted(glob.glob(os.path.join(carpeta, "pag*.png")))
    partes = []
    for pagina in paginas[:MAX_PAGINAS_OCR]:
        r = ejecutar(["tesseract", pagina, "stdout", "-l", IDIOMA_OCR], timeout=300)
        partes.append(r.stdout.decode("utf-8", "replace"))

    return "\n\n".join(partes), f"OCR de {len(paginas)} página(s) escaneada(s)"


def texto_de_imagen(ruta):
    r = ejecutar(["tesseract", ruta, "stdout", "-l", IDIOMA_OCR], timeout=300)
    return r.stdout.decode("utf-8", "replace"), "OCR de imagen"


def texto_de_xlsx(ruta):
    """Todas las hojas, con sus encabezados: un plan suele venir en varias."""
    from openpyxl import load_workbook

    libro = load_workbook(ruta, data_only=True, read_only=True)
    partes = []
    for hoja in libro.worksheets:
        partes.append(f"### Hoja: {hoja.title}")
        for fila in hoja.iter_rows(values_only=True):
            celdas = [str(c).strip() for c in fila if c is not None and str(c).strip() != ""]
            if celdas:
                partes.append(" | ".join(celdas))
    libro.close()
    return "\n".join(partes), "lectura de hoja de cálculo"


def texto_de_docx(ruta):
    import docx

    documento = docx.Document(ruta)
    partes = [p.text for p in documento.paragraphs if p.text.strip()]
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                partes.append(" | ".join(celdas))
    return "\n".join(partes), "lectura de documento Word"


def texto_por_libreoffice(ruta, carpeta, extension):
    """Recurso para formatos antiguos (.doc, .xls, .odt) que no leen las otras."""
    destino = "csv" if extension in ("xls", "ods") else "txt:Text"
    ejecutar([
        "libreoffice", "--headless", "--norestore",
        "--convert-to", destino, "--outdir", carpeta, ruta,
    ], timeout=300, cwd=carpeta)

    for candidato in glob.glob(os.path.join(carpeta, "*.txt")) + glob.glob(os.path.join(carpeta, "*.csv")):
        with open(candidato, "r", encoding="utf-8", errors="replace") as fh:
            return fh.read(), "conversión con LibreOffice"

    return "", "conversión con LibreOffice"


@app.get("/salud")
def salud():
    def hay(binario):
        return shutil.which(binario) is not None

    return jsonify({
        "ok": True,
        "tesseract": hay("tesseract"),
        "pdftotext": hay("pdftotext"),
        "pdftoppm": hay("pdftoppm"),
        "libreoffice": hay("libreoffice"),
        "idioma_ocr": IDIOMA_OCR,
        "max_paginas_ocr": MAX_PAGINAS_OCR,
    })


@app.post("/extraer")
def extraer():
    if TOKEN and request.headers.get("X-Token") != TOKEN:
        return jsonify({"ok": False, "error": "Token inválido"}), 403

    archivo = request.files.get("archivo")
    if archivo is None or archivo.filename == "":
        return jsonify({"ok": False, "error": "No llegó ningún archivo"}), 400

    nombre = os.path.basename(archivo.filename)
    extension = nombre.rsplit(".", 1)[-1].lower() if "." in nombre else ""

    carpeta = tempfile.mkdtemp(prefix="doc_")
    try:
        ruta = os.path.join(carpeta, nombre)
        archivo.save(ruta)
        tamano = os.path.getsize(ruta)

        if extension == "pdf":
            texto, metodo = texto_de_pdf(ruta, carpeta)
        elif extension in ("png", "jpg", "jpeg", "webp", "tif", "tiff", "bmp"):
            texto, metodo = texto_de_imagen(ruta)
        elif extension == "xlsx":
            texto, metodo = texto_de_xlsx(ruta)
        elif extension == "docx":
            texto, metodo = texto_de_docx(ruta)
        elif extension in ("xls", "doc", "odt", "ods", "rtf"):
            texto, metodo = texto_por_libreoffice(ruta, carpeta, extension)
        elif extension in ("txt", "csv", "md", "json", "tsv"):
            with open(ruta, "r", encoding="utf-8", errors="replace") as fh:
                texto, metodo = fh.read(), "texto plano"
        else:
            return jsonify({
                "ok": False,
                "error": f"Formato .{extension} no admitido",
            }), 415

        texto = "\n".join(l.rstrip() for l in texto.splitlines())
        while "\n\n\n" in texto:
            texto = texto.replace("\n\n\n", "\n\n")
        texto = texto.strip()

        if not texto:
            return jsonify({
                "ok": False,
                "error": "No se pudo extraer texto. Si es una foto, revise que se lea bien y esté derecha.",
            }), 422

        return jsonify({
            "ok": True,
            "nombre": nombre,
            "bytes": tamano,
            "metodo": metodo,
            "caracteres": len(texto),
            "texto": texto[:200000],
        })

    except subprocess.TimeoutExpired:
        return jsonify({"ok": False, "error": "El documento tardó demasiado en procesarse"}), 504
    except Exception as e:  # noqa: BLE001
        return jsonify({"ok": False, "error": f"Error al procesar: {e}"}), 500
    finally:
        shutil.rmtree(carpeta, ignore_errors=True)
